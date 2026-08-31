"""
Word (.docx) Document Parser using python-docx.
"""

from __future__ import annotations
import io
import os
from typing import Any, Dict, List

from .base import BaseDocumentParser, DocumentSection, ParsedDocument

class DocxParser(BaseDocumentParser):
    """Parses Microsoft Word (.docx) documents."""

    def parse_bytes(self, file_bytes: bytes, filename: str = "document.docx") -> ParsedDocument:
        try:
            import docx
            doc = docx.Document(io.BytesIO(file_bytes))
            sections: List[DocumentSection] = []
            paragraphs_text: List[str] = []
            tables_data: List[List[List[str]]] = []

            current_section_title = "Document Content"
            current_section_lines: List[str] = []

            for p in doc.paragraphs:
                text = p.text.strip()
                if not text:
                    continue

                if p.style.name.startswith("Heading"):
                    try:
                        level = int(p.style.name.replace("Heading", "").strip())
                    except ValueError:
                        level = 1

                    if current_section_lines:
                        sections.append(
                            DocumentSection(
                                title=current_section_title,
                                content="\n".join(current_section_lines),
                                level=1,
                            )
                        )
                        current_section_lines = []

                    current_section_title = text

                paragraphs_text.append(text)
                current_section_lines.append(text)

            if current_section_lines:
                sections.append(
                    DocumentSection(
                        title=current_section_title,
                        content="\n".join(current_section_lines),
                        level=1,
                    )
                )

            # Extract tables
            for table in doc.tables:
                table_matrix: List[List[str]] = []
                for row in table.rows:
                    row_data = [cell.text.strip() for cell in row.cells]
                    table_matrix.append(row_data)
                if table_matrix:
                    tables_data.append(table_matrix)

            raw_text = "\n\n".join(paragraphs_text)

            # Metadata
            meta: Dict[str, Any] = {}
            if hasattr(doc, "core_properties"):
                props = doc.core_properties
                meta = {
                    "title": props.title or "",
                    "author": props.author or "",
                    "created": str(props.created) if props.created else "",
                    "modified": str(props.modified) if props.modified else "",
                }

            return ParsedDocument(
                filename=filename,
                file_type="docx",
                raw_text=raw_text,
                sections=sections,
                tables=tables_data,
                metadata=meta,
                page_count=max(1, len(paragraphs_text) // 15),
            )
        except Exception as ex:
            raise ValueError(f"Failed to parse DOCX document '{filename}': {str(ex)}")

    def parse_file(self, file_path: str) -> ParsedDocument:
        if not os.path.exists(file_path):
            raise FileNotFoundError(f"DOCX file not found: {file_path}")
        with open(file_path, "rb") as f:
            return self.parse_bytes(f.read(), filename=os.path.basename(file_path))
