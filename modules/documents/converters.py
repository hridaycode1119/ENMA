"""
Multi-Format Document Conversion Engine.
Exports document text into styled PDF, Word DOCX, Markdown, and TXT files.
"""

from __future__ import annotations
import io
import re
from typing import Optional

class DocumentConverter:
    """Converts structured text or markdown into downloadable binary documents."""

    @classmethod
    def text_to_pdf(cls, text: str, title: str = "Document") -> bytes:
        """Renders text or markdown into a clean, formatted PDF using ReportLab."""
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
            from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
            from reportlab.lib import colors

            buf = io.BytesIO()
            doc = SimpleDocTemplate(
                buf,
                pagesize=letter,
                rightMargin=54,
                leftMargin=54,
                topMargin=54,
                bottomMargin=54,
            )
            styles = getSampleStyleSheet()

            # Custom typography styles
            title_style = ParagraphStyle(
                "DocTitle",
                parent=styles["Heading1"],
                fontSize=20,
                leading=24,
                textColor=colors.HexColor("#1E293B"),
                spaceAfter=14,
            )
            h2_style = ParagraphStyle(
                "DocH2",
                parent=styles["Heading2"],
                fontSize=14,
                leading=18,
                textColor=colors.HexColor("#2563EB"),
                spaceBefore=12,
                spaceAfter=6,
            )
            body_style = ParagraphStyle(
                "DocBody",
                parent=styles["Normal"],
                fontSize=10.5,
                leading=15,
                textColor=colors.HexColor("#334155"),
                spaceAfter=8,
            )

            story = []
            if title:
                story.append(Paragraph(title, title_style))
                story.append(Spacer(1, 10))

            lines = text.splitlines()
            current_p = []

            for line in lines:
                stripped = line.strip()
                if not stripped:
                    if current_p:
                        clean_text = " ".join(current_p).replace("<", "&lt;").replace(">", "&gt;")
                        story.append(Paragraph(clean_text, body_style))
                        current_p = []
                    continue

                if stripped.startswith("# "):
                    if current_p:
                        story.append(Paragraph(" ".join(current_p).replace("<", "&lt;").replace(">", "&gt;"), body_style))
                        current_p = []
                    story.append(Paragraph(stripped[2:], title_style))
                elif stripped.startswith("## ") or stripped.startswith("### "):
                    if current_p:
                        story.append(Paragraph(" ".join(current_p).replace("<", "&lt;").replace(">", "&gt;"), body_style))
                        current_p = []
                    story.append(Paragraph(stripped.lstrip("#").strip(), h2_style))
                else:
                    current_p.append(stripped)

            if current_p:
                clean_text = " ".join(current_p).replace("<", "&lt;").replace(">", "&gt;")
                story.append(Paragraph(clean_text, body_style))

            doc.build(story)
            return buf.getvalue()
        except Exception as ex:
            raise RuntimeError(f"PDF compilation failed: {str(ex)}")

    @classmethod
    def text_to_docx(cls, text: str, title: str = "Document") -> bytes:
        """Renders text or markdown into a Microsoft Word (.docx) binary."""
        try:
            import docx
            doc = docx.Document()

            if title:
                doc.add_heading(title, level=0)

            for line in text.splitlines():
                stripped = line.strip()
                if not stripped:
                    continue

                if stripped.startswith("# "):
                    doc.add_heading(stripped[2:], level=1)
                elif stripped.startswith("## "):
                    doc.add_heading(stripped[3:], level=2)
                elif stripped.startswith("### "):
                    doc.add_heading(stripped[4:], level=3)
                elif stripped.startswith("- ") or stripped.startswith("* "):
                    doc.add_paragraph(stripped[2:], style="List Bullet")
                else:
                    doc.add_paragraph(stripped)

            buf = io.BytesIO()
            doc.save(buf)
            return buf.getvalue()
        except Exception as ex:
            raise RuntimeError(f"DOCX compilation failed: {str(ex)}")

    @classmethod
    def convert(cls, text: str, target_format: str, title: str = "Exported Document") -> bytes:
        """Unified conversion dispatcher."""
        fmt = target_format.lower().lstrip(".")
        if fmt == "pdf":
            return cls.text_to_pdf(text, title=title)
        elif fmt in ("docx", "doc"):
            return cls.text_to_docx(text, title=title)
        elif fmt in ("md", "markdown", "txt", "text"):
            return text.encode("utf-8")
        else:
            raise ValueError(f"Unsupported export format: {target_format}")
